from __future__ import annotations

import csv
import json
import os
import re
import shutil
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook, Workbook
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

ROOT = Path(__file__).resolve().parent
CONFIG_PATH = ROOT / "tpn_automation_config.json"
JSON_PATH = ROOT / "dedicated-day-data.json"

def load_config():
    cfg = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    cfg["download_folder"] = Path(
        os.path.expandvars(os.path.expanduser(cfg.get("download_folder", str(ROOT))))
    ).resolve()
    cfg["browser_profile_folder"] = Path(
        os.path.expandvars(os.path.expanduser(cfg.get("browser_profile_folder", str(ROOT / "browser_profile"))))
    ).resolve()
    return cfg

def parse_iso_dates(values):
    out=set()
    for v in values or []:
        try: out.add(datetime.strptime(v,"%Y-%m-%d").date())
        except ValueError: pass
    return out

def last_working_day(today, holidays):
    d=today-timedelta(days=1)
    while d.weekday()>=5 or d in holidays:
        d-=timedelta(days=1)
    return d

def working_day_range(end_date, count, holidays):
    days=[]
    d=end_date
    while len(days)<count:
        if d.weekday()<5 and d not in holidays:
            days.append(d)
        d-=timedelta(days=1)
    return min(days), max(days)

def fmt_ui(d):
    return d.strftime("%d/%m/%Y")

def safe_click(locator, timeout=5000):
    try:
        locator.first.click(timeout=timeout)
        return True
    except Exception:
        return False

def click_exact_nav(page, name):
    candidates=[
        page.get_by_role("link",name=name,exact=True),
        page.get_by_role("button",name=name,exact=True),
        page.locator(f'a:has-text("{name}")').filter(has_text=re.compile(rf"^\s*{re.escape(name)}\s*$",re.I)),
        page.locator(f'button:has-text("{name}")').filter(has_text=re.compile(rf"^\s*{re.escape(name)}\s*$",re.I)),
    ]
    for loc in candidates:
        if safe_click(loc): return
    raise RuntimeError(f"Could not find top-level navigation item {name!r}")

def fill_by_label_or_placeholder(page, names, value):
    for name in names:
        for loc in (page.get_by_label(re.compile(name,re.I)), page.get_by_placeholder(re.compile(name,re.I))):
            try:
                if loc.count():
                    loc.first.fill(value,timeout=3000)
                    return True
            except Exception:
                pass
    return False

def set_date_range(page, start_date, end_date):
    fv,tv=fmt_ui(start_date),fmt_ui(end_date)
    okf=fill_by_label_or_placeholder(page,[r"date\s*from",r"\bfrom\b",r"start\s*date"],fv)
    okt=fill_by_label_or_placeholder(page,[r"date\s*to",r"\bto\b",r"end\s*date"],tv)
    if okf and okt: return
    inputs=page.locator('input:visible[type="date"], input:visible.k-input, input:visible[type="text"]')
    if inputs.count()>=2:
        inputs.nth(0).fill(fv)
        inputs.nth(1).fill(tv)
        return
    raise RuntimeError("Could not identify Date From / Date To controls")

def wait_for_manual_login(page,cfg):
    print("Log in to TPN if requested. Waiting for the main Browse menu...")
    timeout_ms=int(cfg.get("manual_login_timeout_minutes",10))*60*1000
    try:
        page.get_by_text("Browse",exact=True).first.wait_for(state="visible",timeout=timeout_ms)
    except PlaywrightTimeoutError:
        raise RuntimeError("Timed out waiting for TPN login")

def browse_export(page,cfg,start_date,end_date):
    click_exact_nav(page,"Browse")
    page.wait_for_timeout(1000)
    set_date_range(page,start_date,end_date)

    btn=page.get_by_role("button",name="Browse",exact=True)
    if btn.count(): btn.first.click()
    else: page.get_by_text("Browse",exact=True).last.click()
    page.wait_for_timeout(1500)

    stamp=datetime.now().strftime("%Y%m%d-%H%M%S")
    dest=cfg["download_folder"]/f"RAW_TPN Dedicated Day Status Refresh({stamp}).xlsx"
    with page.expect_download(timeout=120000) as info:
        x=page.get_by_role("button",name=re.compile(r"Export\s*To\s*Excel",re.I))
        if x.count(): x.first.click()
        else:
            x=page.get_by_text(re.compile(r"Export\s*To\s*Excel",re.I))
            if not x.count(): raise RuntimeError("Could not find Export To Excel")
            x.first.click()
    info.value.save_as(dest)
    print("Downloaded:",dest.name)
    return dest

def norm_header(v):
    return re.sub(r"\s+"," ",str(v or "").strip()).lower()

def to_number(v):
    try:
        if v is None or str(v).strip()=="": return None
        return float(str(v).strip())
    except Exception:
        return None

def locally_filter_browse(xlsx_path,cfg):
    wb=load_workbook(xlsx_path,read_only=True,data_only=True)
    ws=wb.active
    it=ws.iter_rows(values_only=True)
    headers=list(next(it))
    rows=list(it)

    idx={}
    for wanted in ("Docket","Service","Req","Del","STATUS"):
        exact=[i for i,h in enumerate(headers) if norm_header(h)==norm_header(wanted)]
        idx[wanted]=exact[0] if exact else None

    missing=[k for k,v in idx.items() if v is None]
    if missing:
        raise RuntimeError("Browse export is missing required columns: "+", ".join(missing))

    prefix=str(cfg.get("service_prefix","DD")).upper()
    kept=[]
    for r in rows:
        service=str(r[idx["Service"]] or "").strip().upper()
        req=to_number(r[idx["Req"]])
        dele=to_number(r[idx["Del"]])
        if service.startswith(prefix) and req==8 and dele!=8:
            kept.append(r)

    filtered=xlsx_path.with_name(xlsx_path.stem+"_FILTERED.xlsx")
    out=Workbook(write_only=True)
    ws2=out.create_sheet("Filtered")
    ws2.append(headers)
    for r in kept: ws2.append(list(r))
    out.save(filtered)

    print(f"Browse rows: {len(rows)} -> {len(kept)} after local filters")
    return filtered

def norm_docket(v):
    s=str(v or "").strip()
    if s.endswith(".0"): s=s[:-2]
    return re.sub(r"\s+","",s).upper()

def update_statuses_from_browse(filtered_xlsx):
    if not JSON_PATH.exists():
        raise RuntimeError("dedicated-day-data.json was not found. Run the morning full refresh first.")

    wb=load_workbook(filtered_xlsx,read_only=True,data_only=True)
    ws=wb.active
    it=ws.iter_rows(values_only=True)
    headers=list(next(it))
    hm={norm_header(h):i for i,h in enumerate(headers)}
    di=hm.get("docket")
    si=hm.get("status")
    if di is None or si is None:
        raise RuntimeError("Filtered Browse file does not contain Docket and STATUS")

    status_lookup={}
    for r in it:
        d=norm_docket(r[di])
        if d:
            status_lookup[d]=str(r[si] or "").strip()

    payload=json.loads(JSON_PATH.read_text(encoding="utf-8"))
    rows=payload.get("rows",[])
    matched=0
    changed=0
    before={}
    for row in rows:
        d=norm_docket(row.get("Docket"))
        old=str(row.get("Status","") or "")
        before[d]=old
        if d in status_lookup:
            matched+=1
            new=status_lookup[d]
            if new!=old:
                changed+=1
                row["Status"]=new

    payload["generated_at"]=datetime.now().astimezone().isoformat(timespec="seconds")
    payload.setdefault("source_files",{})["status_refresh_browse"]=filtered_xlsx.name
    payload["status_refresh"]={
        "refreshed_at":datetime.now().astimezone().isoformat(timespec="seconds"),
        "matched_dockets":matched,
        "changed_statuses":changed
    }

    tmp=JSON_PATH.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    tmp.replace(JSON_PATH)

    print(f"Dashboard statuses refreshed: {matched} dockets matched; {changed} status values changed.")
    return matched,changed

def main():
    cfg=load_config()
    cfg["download_folder"].mkdir(parents=True,exist_ok=True)
    cfg["browser_profile_folder"].mkdir(parents=True,exist_ok=True)

    holidays=parse_iso_dates(cfg.get("bank_holidays",[]))
    mode=str(cfg.get("date_mode","last_working_day")).lower()
    if mode=="today":
        end=date.today()
    elif mode.startswith("fixed:"):
        end=datetime.strptime(mode.split(":",1)[1],"%Y-%m-%d").date()
    else:
        end=last_working_day(date.today(),holidays)

    start,end=working_day_range(end,int(cfg.get("working_days_to_export",7)),holidays)
    print(f"Browse status refresh range: {fmt_ui(start)} to {fmt_ui(end)}")
    print("Local filters: Service starts DD; Req=8; Del!=8")

    with sync_playwright() as p:
        context=p.chromium.launch_persistent_context(
            user_data_dir=str(cfg["browser_profile_folder"]),
            headless=False,
            accept_downloads=True,
            viewport={"width":1500,"height":950},
        )
        page=context.pages[0] if context.pages else context.new_page()
        page.goto(cfg.get("base_url","https://connect.tpnsecure.com/"),wait_until="domcontentloaded",timeout=120000)
        wait_for_manual_login(page,cfg)
        raw=browse_export(page,cfg,start,end)
        context.close()

    filtered=locally_filter_browse(raw,cfg)
    update_statuses_from_browse(filtered)
    print("Status refresh complete.")

if __name__=="__main__":
    try:
        main()
    except Exception as exc:
        print("\nERROR:",exc)
        raise SystemExit(1)
